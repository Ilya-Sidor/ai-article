"""Case data from documents (FR-1.1): .docx, .doc, .pdf, .rtf, .txt.

A document is either a table of cases (a spreadsheet pasted into Word, a PDF export of a register) or a
free-text report (a pathology report per case). A table becomes rows as usual; a report becomes one case
with its text in DOC_TEXT_COLUMN, anonymised like any text cell and turned into features later by
``case_extraction`` after the author confirms the anonymisation. Only text and table cells are read, so
document properties (author, company, revision history) never enter the project (FR-1.6).
"""
import io
import shutil
import subprocess
import tempfile
from pathlib import Path

import pandas as pd

DOCUMENT_EXT = {".docx", ".doc", ".pdf", ".rtf", ".txt"}
DOC_TEXT_COLUMN = "Текст документа"
TABLE_SHARE = 0.6  # a document is a case table when its tables hold most of its text


class DocumentError(ValueError):
    pass


def read_document(ext: str, content: bytes) -> pd.DataFrame:
    if ext in (".doc", ".rtf"):
        content, ext = _convert_to_docx(ext, content), ".docx"
    if ext == ".docx":
        text, tables = _docx(content)
    elif ext == ".pdf":
        text, tables = _pdf(content)
    else:
        text, tables = _decode(content), []
    frame = _case_table(text, tables)
    if frame is not None:
        return frame
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < 20:
        raise DocumentError("в документе не найден текст — если это скан, распознайте его (OCR) и загрузите снова")
    return pd.DataFrame({DOC_TEXT_COLUMN: [text]})


def _decode(content):
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    raise DocumentError("не удалось определить кодировку текста")


def _convert_to_docx(ext, content):
    """Old Word (.doc) and RTF via macOS textutil or LibreOffice, whichever is installed."""
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / f"in{ext}"
        src.write_bytes(content)
        if shutil.which("textutil"):
            cmd = ["textutil", "-convert", "docx", "-output", str(Path(tmp) / "in.docx"), str(src)]
        elif shutil.which("soffice"):
            cmd = ["soffice", "--headless", "--convert-to", "docx", "--outdir", tmp, str(src)]
        else:
            raise DocumentError(f"формат {ext} не поддерживается на этом сервере — сохраните файл как .docx")
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=120)
            return (Path(tmp) / "in.docx").read_bytes()
        except (subprocess.SubprocessError, OSError) as exc:
            raise DocumentError(f"не удалось прочитать {ext} — сохраните файл как .docx") from exc


def _cells(rows):
    return [[" ".join((c or "").split()) for c in row] for row in rows]


def _docx(content):
    from docx import Document
    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        raise DocumentError("не удалось прочитать .docx") from exc
    tables = [_cells([[cell.text for cell in row.cells] for row in t.rows]) for t in doc.tables]
    # tables stay part of the text: an IHC panel inside a report is data too
    flat = ["\n".join(" | ".join(c for c in row if c) for row in t) for t in tables]
    return "\n".join([p.text for p in doc.paragraphs] + flat), tables


def _pdf(content):
    import pdfplumber
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = [(p.extract_text() or "", [_cells(t) for t in p.extract_tables() if t]) for p in pdf.pages]
    except Exception as exc:
        raise DocumentError("не удалось прочитать PDF") from exc
    return "\n".join(t for t, _ in pages), [t for _, ts in pages for t in ts]


def _case_table(text, tables):
    """Rows of cases when the document is essentially one table (possibly split over pages)."""
    tables = [[r for r in t if any(r)] for t in tables]
    tables = [t for t in tables if len(t) >= 2 and len(t[0]) >= 2]
    if not tables:
        return None
    main = max(tables, key=lambda t: len(t) * len(t[0]))
    header, width = main[0], len(main[0])
    rows = []
    for t in tables:  # the same table continued on the next pages, with or without a repeated header
        if len(t[0]) != width:
            continue
        rows += t[1:] if (t is main or t[0] == header) else t
    table_chars = sum(len(c) for r in [header] + rows for c in r)
    outside = len("".join(text.split())) - sum(len("".join(c.split())) for r in [header] + rows for c in r)
    if len(rows) < 2 or table_chars < TABLE_SHARE * (table_chars + max(outside, 0)):
        return None  # a report with a small embedded table (e.g. an IHC panel) stays a report
    names, seen = [], {}
    for i, h in enumerate(header):
        h = h or f"Столбец {i + 1}"
        seen[h] = seen.get(h, 0) + 1
        names.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    return pd.DataFrame(rows, columns=names)
