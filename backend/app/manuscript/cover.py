"""Cover letter (Module 6: FR-6.1 – FR-6.6).

The model writes only the body of the letter. The salutation, the reviewer lists
and the signature block are assembled by code from the author's settings, so
names and e-mail addresses of reviewers and authors are never sent to the LLM.
"""
import difflib
import io
import re

from .. import llm
from ..config import MODEL_REASONING
from ..storage import now_iso
from .agent import FALLBACK

TONES = {"formal": "formal and concise", "personal": "warm but professional, first person plural, slightly more personal"}

REQUIRED = [
    ("originality", r"\b(original|not been published|has not been published|unpublished)\b",
     "заявление об оригинальности (работа не публиковалась)"),
    ("exclusive", r"\b(not under consideration|not being considered|not been submitted|not submitted elsewhere|"
                  r"exclusively)\b", "заявление, что рукопись не подана в другой журнал"),
    ("coi", r"\b(conflicts? of interest|competing interests?)\b", "заявление о конфликте интересов"),
    ("approval", r"\ball (the )?authors (have )?(read and )?approved\b", "все авторы одобрили рукопись"),
]
HYPE = re.compile(r"\b(groundbreaking|ground-breaking|first[- ]ever|unprecedented|paradigm[- ]shift\w*|"
                  r"revolutionar\w+|landmark|definitive(ly)?|prove[sn]?)\b", re.I)

SCHEMA = {
    "type": "object",
    "properties": {
        "body": {"type": "string"},
        "requirements_coverage": {"type": "array", "items": {
            "type": "object",
            "properties": {"requirement": {"type": "string"}, "addressed": {"type": "boolean"},
                           "where": {"type": "string"}},
            "required": ["requirement", "addressed", "where"], "additionalProperties": False}},
        "questions": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["body", "requirements_coverage", "questions"], "additionalProperties": False,
}

SYSTEM = """You write the body of a cover letter for submitting a pathology manuscript to a specific journal.

Content, in this order: the manuscript title and article type; why the work fits this journal's scope and
readers (use the journal scope/tone given, be specific); the key findings and what they add, stated with the
strength their evidence level allows (exploratory findings are observations in this series that warrant
validation — never "prove", "establish", "first ever", "groundbreaking"); then the declarations: the work is
original and has not been published, it is not under consideration elsewhere, all authors approved the
manuscript, and the conflict-of-interest statement given by the authors. Address every journal-specific
requirement listed and report in requirements_coverage where you addressed it (quote the phrase).

Rules:
- Do NOT write the salutation, the signature, author names, e-mails, or reviewer names: they are added
  automatically after your text. Start with the first paragraph and end with the closing sentence
  (e.g. "Thank you for considering our manuscript.").
- Numbers from the study must be written as fact placeholders ({{n}}, {{A1.p_expr}}); prefer no p-values in a
  cover letter.
- If something needed is missing, write [уточнить: <what, in Russian>].
- 250–400 words. English ({variant}). Tone: {tone}.
- questions: what you need from the authors (in Russian)."""


def empty():
    return {"settings": {"tone": "formal"}, "versions": [], "status": "empty", "comments": [], "confirmed": {}}


def current(cover):
    return cover["versions"][-1] if cover["versions"] else None


def add_version(cover, body, actor, note, extra=None):
    cur = current(cover)
    if cur and cur["body"] == body and actor == "author":
        return cur
    v = {"n": len(cover["versions"]) + 1, "ts": now_iso(), "actor": actor, "note": note, "body": body}
    if extra:
        v.update(extra)
    cover["versions"].append(v)
    cover["status"] = "draft"
    return v


def diff(cover, a, b):
    va = next(v for v in cover["versions"] if v["n"] == a)
    vb = next(v for v in cover["versions"] if v["n"] == b)
    split = lambda t: re.split(r"(?<=[.!?])\s+", t)  # noqa: E731
    return list(difflib.unified_diff(split(va["body"]), split(vb["body"]), f"v{a}", f"v{b}", n=1, lineterm=""))


def _context(ctx, cover, front, abstract_text, statements_source):
    t = ctx.get("template") or {}
    s = cover["settings"]
    coi = re.search(r"### Conflict of interest\s+(.+?)(?:\n### |\Z)", statements_source or "", re.S)
    return {
        "journal": ctx.get("journal"), "journal_scope": (ctx.get("profile") or {}).get("fields", {}).get("scope", {})
        .get("value") if ctx.get("profile") else None,
        "journal_tone": t.get("tone"), "article_type": t.get("type_name") or ctx.get("article_type"),
        "journal_cover_letter_requirements": (t.get("cover_letter") or {}).get("requirements") or [],
        "title": front.get("title"), "key_messages": ctx.get("key_messages"),
        "confirmed_novelty": ctx.get("novelty"), "findings": ctx.get("findings"), "abstract": abstract_text,
        "conflict_of_interest_statement": coi.group(1).strip() if coi else None,
        "transfer_or_resubmission": s.get("previous_submission") or None,
        "suggested_reviewers_listed": bool(s.get("suggested_reviewers")),
        "excluded_reviewers_listed": bool(s.get("excluded_reviewers")),
    }


def _facts(ctx):
    return "\n".join(f"{k} — {v['desc']} = {v['value']}" for k, v in ctx["facts"].items() if not k.startswith("V"))


def generate(project, ctx, cover, front, abstract_text, statements_source):
    import json
    s = cover["settings"]
    system = SYSTEM.format(variant=ctx.get("language_variant") or "consistent spelling",
                           tone=TONES.get(s.get("tone"), TONES["formal"]))
    user = (json.dumps(_context(ctx, cover, front, abstract_text, statements_source), ensure_ascii=False, indent=1)
            + "\n\nFacts (use ids as placeholders):\n" + _facts(ctx))
    return llm._call(project, "cover_letter", MODEL_REASONING, system, user, SCHEMA, max_tokens=16000, gate=True,
                     output_config={"effort": "medium"}, **FALLBACK)


def revise(project, ctx, cover, body, instruction, selection, front, abstract_text, statements_source):
    import json
    s = cover["settings"]
    system = SYSTEM.format(variant=ctx.get("language_variant") or "consistent spelling",
                           tone=TONES.get(s.get("tone"), TONES["formal"]))
    user = (f"Revise the cover letter body according to the authors' instruction (may be in Russian): {instruction}\n"
            + (f"Apply it only to this fragment, keep the rest: «{selection}»\n" if selection else "")
            + f"Current body:\n{body}\n\nContext:\n"
            + json.dumps(_context(ctx, cover, front, abstract_text, statements_source), ensure_ascii=False, indent=1)
            + "\n\nFacts:\n" + _facts(ctx))
    return llm._call(project, "revise_cover_letter", MODEL_REASONING, system, user, SCHEMA, max_tokens=16000,
                     gate=True, output_config={"effort": "medium"}, **FALLBACK)


# ---------------------------------------------------------------------------
# Assembly by code: salutation, reviewers, signature
# ---------------------------------------------------------------------------

def _people(text):
    return [line.strip() for line in (text or "").splitlines() if line.strip()]


def assemble(cover, rendered_body, journal):
    """Full letter as a list of (kind, text) blocks."""
    s = cover["settings"]
    editor = (s.get("editor") or "").strip()
    blocks = [("date", now_iso()[:10]), ("to", f"{editor + ', ' if editor else ''}Editor-in-Chief\n{journal or ''}".strip()),
              ("p", (s.get("salutation") or "").strip()
               or f"Dear {('Dr ' + editor.split()[-1]) if editor else 'Editor'},")]
    blocks += [("p", p.strip()) for p in re.split(r"\n\s*\n", rendered_body.strip()) if p.strip()]
    sugg, excl = _people(s.get("suggested_reviewers")), _people(s.get("excluded_reviewers"))
    if sugg:
        blocks.append(("p", "We would like to suggest the following potential reviewers, who have expertise in this "
                            "field and no conflict of interest with the authors:"))
        blocks += [("li", x) for x in sugg]
    if excl:
        blocks.append(("p", "We respectfully request that the following individuals not be invited to review the "
                            "manuscript:"))
        blocks += [("li", x) for x in excl]
    blocks.append(("p", "Yours sincerely,"))
    sig = [s.get("corresponding_name"), s.get("corresponding_affiliation"), s.get("corresponding_email")]
    blocks.append(("sig", "\n".join(x for x in sig if x) or "[уточнить: подпись и контакты автора для переписки]"))
    return blocks


def as_text(blocks):
    return "\n\n".join(("• " + t) if k == "li" else t for k, t in blocks)


def check(cover, rendered_body, raw_body, front, template, ctx):
    """Deterministic checks of the letter (issues like manuscript checks)."""
    issues = []

    def add(sev, code, message, fragment="", suggestion=""):
        issues.append({"section": "cover_letter", "heading": "Cover letter", "severity": sev, "code": code,
                       "message": message, "fragment": fragment[:200], "suggestion": suggestion})

    text = rendered_body
    for key, rx, label in REQUIRED:
        if not re.search(rx, text, re.I):
            add("blocking", f"missing_{key}", f"нет обязательного элемента: {label}")
    title = (front or {}).get("title")
    if not title:
        add("blocking", "no_title", "в рукописи не выбрано название")
    elif title.lower().strip(" .") not in re.sub(r"\s+", " ", text.lower()):
        add("warning", "title_mismatch", "название статьи в письме не совпадает с выбранным", title,
            "вставьте точное название")
    for m in HYPE.finditer(text):
        add("warning", "hype", f"слишком сильная формулировка: «{m.group(0)}»", m.group(0),
            "в cover letter формулировки не сильнее, чем в статье")
    if re.search(r"\bsignificant(ly)?\b", text, re.I) and any(f["evidence"] != "significant" for f in ctx["findings"]):
        add("warning", "evidence_wording", "«significant» при наличии exploratory-находок — проверьте формулировку")
    for m in re.finditer(r"\[(?:уточнить|TODO)\s*:\s*([^\]]+)\]", raw_body, re.I):
        add("blocking", "todo", f"требуется информация: {m.group(1)}", m.group(0))
    for m in re.finditer(r"\{\{\s*([A-Za-z0-9_.:-]+)\s*\}\}", raw_body):
        if m.group(1) not in ctx["facts"]:
            add("blocking", "unknown_ref", f"неизвестная ссылка {m.group(0)}", m.group(0))
    if re.search(r"^\s*dear\b", raw_body, re.I | re.M) or re.search(r"yours (sincerely|faithfully)", raw_body, re.I):
        add("warning", "duplicate_frame", "обращение/подпись в теле письма — они добавляются автоматически")
    words = len(re.findall(r"\w+", text))
    if words > 500:
        add("warning", "length", f"{words} слов — обычно cover letter ≤ 1 страницы")
    s = cover["settings"]
    if not (s.get("corresponding_name") and s.get("corresponding_email")):
        add("blocking", "signature", "не указан автор для переписки (имя и e-mail)")
    reqs = (template or {}).get("cover_letter", {}).get("requirements") or []
    unconfirmed = [r for r in reqs if not cover.get("confirmed", {}).get(r)]
    if unconfirmed:
        add("blocking", "requirements", f"автор не подтвердил выполнение требований журнала: {len(unconfirmed)}",
            "; ".join(unconfirmed)[:200])
    if (template or {}).get("cover_letter", {}).get("suggested_reviewers") == "required" and not s.get(
            "suggested_reviewers"):
        add("blocking", "reviewers", "журнал требует предложить рецензентов")
    return issues


def docx_bytes(blocks, draft):
    from docx import Document
    from docx.shared import Pt, RGBColor

    from .render import plain_text
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(8)
    if draft:
        r = doc.sections[0].header.paragraphs[0].add_run("DRAFT — NOT FOR SUBMISSION")
        r.bold = True
        r.font.color.rgb = RGBColor(0xB0, 0x10, 0x30)
    for kind, text in blocks:
        text = plain_text(text) or ""
        if kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            p = doc.add_paragraph()
            for i, line in enumerate(text.split("\n")):
                run = p.add_run(line)
                if "[уточнить" in line:
                    run.font.color.rgb = RGBColor(0xB0, 0x10, 0x30)
                if i < len(text.split("\n")) - 1:
                    run.add_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
