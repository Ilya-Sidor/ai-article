"""Response to reviewers (after submission): points, responses, changes, the marked-up manuscript.

A revision round starts from the editor's decision letter: the model splits it into the reviewers' points
(each must be a verbatim part of the letter), drafts a response and the change to the text for each point;
the change is applied as an ordinary section revision. The round keeps a snapshot of the submitted text, so the
manuscript with changes marked (insertions underlined, deletions struck through) and the point-by-point
response letter can be exported.
"""
import difflib
import io
import re

from .. import llm
from ..config import MODEL_EXTRACTION, MODEL_REASONING
from ..storage import now_iso

SPLIT_SCHEMA = {
    "type": "object",
    "properties": {"points": {"type": "array", "items": {
        "type": "object",
        "properties": {"reviewer": {"type": "string"}, "text": {"type": "string"},
                       "kind": {"type": "string", "enum": ["major", "minor", "editorial", "editor"]}},
        "required": ["reviewer", "text", "kind"], "additionalProperties": False}}},
    "required": ["points"], "additionalProperties": False,
}
SPLIT_SYSTEM = """You split an editor's decision letter with reviewers' reports into separate points to answer.
Each point: reviewer ("Editor", "Reviewer 1", "Reviewer 2"… as in the letter), text — the comment copied from
the letter character for character (one request or question per point; keep numbering out), kind — major,
minor, editorial (language/formatting) or editor (a requirement of the editor). Skip greetings, signatures and
general praise that asks for nothing."""

DRAFT_SCHEMA = {
    "type": "object",
    "properties": {"response": {"type": "string"}, "change_needed": {"type": "boolean"},
                   "section": {"type": "string"}, "quote": {"type": "string"}, "instruction": {"type": "string"}},
    "required": ["response", "change_needed", "section", "quote", "instruction"], "additionalProperties": False,
}
DRAFT_SYSTEM = """You help the authors answer one reviewer comment on their pathology manuscript.
Write the response the way experienced authors do: thank briefly, answer the substance, state exactly what was
changed and where (section), or — if the comment cannot or should not be followed — explain politely with
reasons (e.g. data not available in a retrospective series). Never promise analyses or data that the study does
not have; if the authors must supply something (a clone, a number, a reference), write [уточнить: …].
Language of the response: {language}.
change_needed: whether the manuscript text should change. section: the heading of the section to change,
quote: a passage copied character for character from that section where the change goes (or the nearest
sentence), instruction: in Russian, a precise editing instruction for that passage. If no change: empty
section, quote and instruction."""


def _norm(s):
    return re.sub(r"\s+", " ", (s or "").replace("ё", "е")).strip().lower()


def split(project, letter):
    out = llm._call(project, "revision_split", MODEL_EXTRACTION, SPLIT_SYSTEM, letter, SPLIT_SCHEMA, max_tokens=12000)
    norm = _norm(letter)
    points = []
    for i, p in enumerate(out.get("points") or [], 1):
        text = p["text"].strip()
        if not text:
            continue
        points.append({"id": f"P{i}", "reviewer": p["reviewer"].strip() or "Reviewer", "kind": p["kind"],
                       "text": text, "verbatim": _norm(text) in norm, "response": "", "change": None,
                       "applied": [], "status": "open"})
    return points


def draft(project, point, manuscript, lang):
    language = "Russian" if lang == "ru" else "English"
    user = (f"Reviewer comment ({point['reviewer']}, {point['kind']}):\n{point['text']}\n\n"
            f"Manuscript (current text):\n{manuscript}")
    return llm._call(project, "revision_draft", MODEL_REASONING, DRAFT_SYSTEM.format(language=language), user,
                     DRAFT_SCHEMA, max_tokens=6000)


# ---------------------------------------------------------------------------
# Export: response letter and the manuscript with changes marked
# ---------------------------------------------------------------------------

def response_docx(round_, title, journal, lang):
    from docx import Document
    from docx.shared import Pt, RGBColor

    from .render import plain_text as _t
    ru = lang == "ru"
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading("Ответ рецензентам" if ru else "Response to reviewers", level=0)
    doc.add_paragraph(_t(f"{'Рукопись' if ru else 'Manuscript'}: {title or '—'}" + (f" ({journal})" if journal else "")))
    doc.add_paragraph("Благодарим редактора и рецензентов за внимательное прочтение рукописи. Ниже приведены ответы "
                      "на каждое замечание; изменения в тексте выделены в прилагаемой версии рукописи." if ru else
                      "We thank the editor and the reviewers for their careful reading of the manuscript. Our "
                      "point-by-point responses follow; changes are marked in the revised manuscript.")
    current = None
    for p in round_["points"]:
        if p["reviewer"] != current:
            current = p["reviewer"]
            doc.add_heading(_t(current), level=1)
        q = doc.add_paragraph()
        q.add_run(f"{p['id']}. ").bold = True
        q.add_run(_t(p["text"])).italic = True
        r = doc.add_paragraph()
        lab = r.add_run(("Ответ: " if ru else "Response: "))
        lab.bold = True
        lab.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)
        r.add_run(_t(p.get("response") or ("[уточнить: ответ]")))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _tokens(text):
    return re.findall(r"\s+|[^\s]+", text)


def marked_docx(sections, baseline, renderer, lang):
    """Each section: the submitted text vs the current one, word by word (inserted underlined blue, deleted
    struck through red). sections: [(key, heading, kind, source)]; baseline: {key: submitted source}."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    from .render import plain_text as _t
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading("Рукопись с выделенными изменениями" if lang == "ru" else "Revised manuscript (changes marked)",
                    level=0)
    for key, heading, kind, source in sections:
        old = renderer.text(baseline.get(key, ""), kind) if baseline.get(key) else ""
        new = renderer.text(source, kind) if source else ""
        if not old and not new:
            continue
        doc.add_heading(_t(heading), level=1)
        for para_old, para_new in _paragraph_pairs(old, new):
            p = doc.add_paragraph()
            a, b = _tokens(para_old), _tokens(para_new)
            for op, i1, i2, j1, j2 in difflib.SequenceMatcher(a=a, b=b, autojunk=False).get_opcodes():
                if op == "equal":
                    p.add_run(_t("".join(a[i1:i2])))
                    continue
                if op in ("delete", "replace"):
                    run = p.add_run(_t("".join(a[i1:i2])))
                    run.font.strike = True
                    run.font.color.rgb = RGBColor(0xB0, 0x10, 0x30)
                if op in ("insert", "replace"):
                    run = p.add_run(_t("".join(b[j1:j2])))
                    run.font.underline = True
                    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _paragraph_pairs(old, new):
    """Align paragraphs of two versions (unchanged, changed, added, removed)."""
    po, pn = old.split("\n\n") if old else [], new.split("\n\n") if new else []
    out = []
    for op, i1, i2, j1, j2 in difflib.SequenceMatcher(a=po, b=pn, autojunk=False).get_opcodes():
        if op == "equal":
            out += [(x, x) for x in po[i1:i2]]
        elif op == "replace":
            a, b = po[i1:i2], pn[j1:j2]
            for k in range(max(len(a), len(b))):
                out.append((a[k] if k < len(a) else "", b[k] if k < len(b) else ""))
        elif op == "delete":
            out += [(x, "") for x in po[i1:i2]]
        else:
            out += [("", x) for x in pn[j1:j2]]
    return out


def new_round(state, points, letter):
    rounds = state.setdefault("revision", {"rounds": []})["rounds"]
    rnd = {"id": f"RV{len(rounds) + 1}", "created_at": now_iso(), "letter": letter, "points": points,
           "baseline": {k: (s["versions"][-1]["source"] if s["versions"] else "") for k, s in state["sections"].items()},
           "front_baseline": dict(state.get("front") or {})}
    rounds.append(rnd)
    return rnd
