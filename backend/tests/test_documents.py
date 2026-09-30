"""Case data from documents (.docx, .doc, .pdf, .txt) and features from report texts (FR-1.1, FR-1.7, FR-1.9)."""
import io
import shutil
import subprocess

import pytest

REPORT_1 = ("Выписной эпикриз. Пациент: Иванов Иван Иванович, 1961 г.р. История болезни № 12345/2024.\n"
            "Поступил 10.03.2024. Операция 12.03.2024: резекция желудка. Рецидив выявлен 20.09.2025.\n"
            "Макроскопически: в стенке желудка опухолевый узел размером 42 мм.\n"
            "Микроскопически: веретеноклеточная гастроинтестинальная стромальная опухоль, митозы 6 на 5 мм², "
            "некроза нет.\nИГХ:")
REPORT_2 = ("Пациентка: Петрова Анна Сергеевна.\nМакроскопически: опухоль тонкой кишки размером 18 мм.\n"
            "Микроскопически: эпителиоидная ГИСО, митозы 1 на 5 мм², очаговый некроз.\nИГХ:")


def _docx(paragraphs, table=None):
    import docx
    d = docx.Document()
    d.core_properties.author = "Доктор Секретов"  # document metadata must never reach the project (FR-1.6)
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, v in enumerate(row):
                t.cell(i, j).text = v
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(lines):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    from matplotlib import font_manager
    font = font_manager.findfont("DejaVu Sans")
    pdfmetrics.registerFont(TTFont("DejaVu", font))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("DejaVu", 10)
    y = 800
    for line in lines:
        c.drawString(40, y, line)
        y -= 16
    c.save()
    return buf.getvalue()


def test_docx_case_table_and_report_with_ihc_panel():
    from app.documents import DOC_TEXT_COLUMN
    from app.ingest import read_table
    table = [["ID", "Возраст", "Локализация", "CD117"], ["1", "54", "желудок", "+"], ["2", "61", "тонкая кишка", "+"],
             ["3", "47", "желудок", "-"]]
    df = read_table("серия.docx", _docx(["Таблица 1. Серия"], table))
    assert list(df.columns) == ["ID", "Возраст", "Локализация", "CD117"] and len(df) == 3
    assert df.iloc[1]["Локализация"] == "тонкая кишка"

    report = _docx(REPORT_1.split("\n"), [["Маркер", "Результат"], ["CD117", "положительный"], ["DOG1", "положительный"]])
    df = read_table("заключение.docx", report)
    assert list(df.columns) == [DOC_TEXT_COLUMN] and len(df) == 1
    text = df.iloc[0][DOC_TEXT_COLUMN]
    assert "42 мм" in text and "CD117 | положительный" in text and "Секретов" not in text


def test_pdf_and_txt_reports():
    from app.documents import DOC_TEXT_COLUMN
    from app.ingest import read_table
    df = read_table("report.pdf", _pdf(REPORT_2.split("\n")))
    assert "опухоль тонкой кишки размером 18 мм" in df.iloc[0][DOC_TEXT_COLUMN]
    df = read_table("report.txt", REPORT_2.replace("²", "2").encode("cp1251"))
    assert "эпителиоидная ГИСО" in df.iloc[0][DOC_TEXT_COLUMN]
    from app.ingest import UnsupportedFile
    with pytest.raises(UnsupportedFile, match="OCR"):
        read_table("scan.pdf", _pdf([]))


@pytest.mark.skipif(not shutil.which("textutil"), reason="macOS textutil converts .doc")
def test_old_word_doc(tmp_path):
    from app.documents import DOC_TEXT_COLUMN
    from app.ingest import read_table
    (tmp_path / "r.txt").write_text(REPORT_1, encoding="utf-8")
    subprocess.run(["textutil", "-convert", "doc", "-output", str(tmp_path / "r.doc"), str(tmp_path / "r.txt")],
                   check=True)
    df = read_table("r.doc", (tmp_path / "r.doc").read_bytes())
    assert "42 мм" in df.iloc[0][DOC_TEXT_COLUMN]


def test_reports_upload_anonymise_and_extract(api, monkeypatch):
    from app import llm
    pid = api.post("/api/projects", json={"title": "Отчёты"}).json()["id"]
    files = [("files", ("Иванов Иван Иванович.docx", _docx(REPORT_1.split("\n")))),
             ("files", ("Петрова А.С.docx", _docx(REPORT_2.split("\n")))),
             ("files", ("Смирнов.pdf", _pdf(["Пациент: Смирнов Олег Петрович.", "Опухоль прямой кишки 30 мм,",
                                            "митозы 12 на 5 мм², некроз есть."])))]
    reports = api.post(f"/api/projects/{pid}/uploads", files=files).json()
    blob = str(reports)
    for secret in ("Иванов", "Петрова", "Смирнов", "12345"):
        assert secret not in blob, secret  # file names and texts are anonymised
    assert api.post(f"/api/projects/{pid}/dataset/extract", json={}).status_code == 400  # not confirmed yet
    api.post(f"/api/projects/{pid}/anonymization/confirm")
    ds = api.get(f"/api/projects/{pid}/dataset").json()
    assert ds["text_columns"] == ["Текст документа"] and len(ds["rows"]) == 3
    texts = {r["case_id"]: r["Текст документа"] for r in ds["rows"]}
    assert all("Иванов" not in t and "Петрова" not in t and "1961" not in t for t in texts.values())
    ivanov = next(c for c, t in texts.items() if "42 мм" in t)
    petrova = next(c for c, t in texts.items() if "18 мм" in t)
    # the timeline survives as intervals from the first date of the document
    assert "[+0.0 мес. от индексной даты]" in texts[ivanov] and "[+18.4 мес. от индексной даты]" in texts[ivanov]
    assert "2024" not in texts[ivanov] and "[ГОД РОЖДЕНИЯ]" in texts[ivanov]
    assert texts[ivanov].startswith("Выписной эпикриз")  # medical words are not taken for names

    calls = []

    def fake_call(project, purpose, model, system, user, schema, **kw):
        calls.append(user)
        assert project["anonymization"]["status"] == "confirmed" and "Иванов" not in user
        return {"features": [{"name": "Размер опухоли, мм", "group": "clinical", "unit": "мм"},
                             {"name": "Митозы на 5 мм²", "group": "morphology", "unit": ""}],
                "cases": [
                    {"case_id": ivanov, "values": [{"feature": "Размер опухоли, мм", "value": "42",
                                                     "quote": "опухолевый узел размером 42 мм"},
                                                    {"feature": "Митозы на 5 мм²", "value": "6",
                                                     "quote": "митозы 6 на 5 мм²"}]},
                    {"case_id": petrova, "values": [{"feature": "Размер опухоли, мм", "value": "18",
                                                     "quote": "размером 18 мм"},
                                                    {"feature": "Митозы на 5 мм²", "value": "3",
                                                     "quote": "митозы 3 на 5 мм²"}]},  # not in the text
                    {"case_id": "C-999", "values": [{"feature": "Размер опухоли, мм", "value": "1", "quote": "x"}]},
                ]}
    monkeypatch.setattr(llm, "_call", fake_call)
    r = api.post(f"/api/projects/{pid}/dataset/extract", json={}).json()
    s = r["summary"]
    assert s["values"] == 4 and s["unverified"] == 1 and s["cases"] == 2
    assert set(s["new_columns"]) == {"Размер опухоли, мм", "Митозы на 5 мм²"}
    rows = {row["case_id"]: row for row in r["rows"]}
    other = next(c for c in rows if c not in (ivanov, petrova))
    assert rows[ivanov]["Размер опухоли, мм"] == "42" and rows[other]["Размер опухоли, мм"] == ""
    assert r["extracted"][ivanov]["Размер опухоли, мм"]["verified"] is True
    assert r["extracted"][petrova]["Митозы на 5 мм²"]["verified"] is False
    dic = {v["name"]: v for v in r["dictionary"]}
    assert dic["Текст документа"]["vtype"] == "text" and "Размер опухоли, мм" in dic

    # the author corrects a value: it is no longer "from AI" and a re-run does not overwrite it
    api.patch(f"/api/projects/{pid}/dataset", json={"case_id": petrova, "column": "Митозы на 5 мм²", "value": "1"})
    r = api.post(f"/api/projects/{pid}/dataset/extract", json={}).json()
    rows = {row["case_id"]: row for row in r["rows"]}
    assert rows[petrova]["Митозы на 5 мм²"] == "1" and "Митозы на 5 мм²" not in r["extracted"][petrova]
    assert r["conflicts"] == [{"case_id": petrova, "column": "Митозы на 5 мм²", "table": "1", "text": "3",
                               "quote": "митозы 3 на 5 мм²"}]
    assert "Размер опухоли, мм" in calls[-1]  # existing columns are offered for consistent names
