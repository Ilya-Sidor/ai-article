"""Response to reviewers: points from the letter, drafted responses, applied changes, marked-up manuscript."""
import io
import zipfile

import pytest
from test_manuscript import Claude, _sec, project  # noqa: F401

LETTER = """Dear Dr Ivanova,
Reviewer 1:
1. Please provide the antibody clone used for CD117.
2. The discussion should compare the results with recent series.
Reviewer 2:
The conclusion overstates the findings.
Sincerely, the Editor"""


class RevClaude(Claude):
    def respond(self, kw):
        system = kw.get("system", "")
        system = system if isinstance(system, str) else " ".join(b["text"] for b in system)
        user = kw["messages"][0]["content"]
        if "split an editor's decision letter" in system:
            return self._r({"points": [
                {"reviewer": "Reviewer 1", "text": "Please provide the antibody clone used for CD117.", "kind": "minor"},
                {"reviewer": "Reviewer 1", "text": "The discussion should compare the results with recent series.", "kind": "major"},
                {"reviewer": "Reviewer 2", "text": "The conclusion overstates findings.", "kind": "major"}]})
        if "answer one reviewer comment" in system:
            return self._r({"response": "We thank the reviewer. The clone was added to Methods.", "change_needed": True,
                            "section": "Materials and methods", "quote": "Immunohistochemistry for CD117",
                            "instruction": "Добавь клон YR145"})
        if "Revise the section" in user:
            return self._r({"text": "Immunohistochemistry for CD117 (clone YR145) was performed.", "new_citations": [],
                            "questions": []})
        return super().respond(kw)


@pytest.fixture()
def claude(monkeypatch):
    from app import llm
    c = RevClaude()
    monkeypatch.setattr(llm, "client", lambda: c)
    return c


def test_revision_round(api, project, claude):  # noqa: F811
    pid, kit, expl = project
    api.post(f"/api/projects/{pid}/manuscript/terms/auto")
    api.post(f"/api/projects/{pid}/manuscript/plan")
    m = api.post(f"/api/projects/{pid}/manuscript/plan/approve").json()
    K = {s["kind"]: s["key"] for s in m["sections"]}
    claude.sections = {"introduction": "GIST is common.", "methods": "Immunohistochemistry for CD117 was performed."}
    for kind in ("introduction", "methods"):
        api.post(f"/api/projects/{pid}/manuscript/sections/{K[kind]}/generate")
        api.post(f"/api/projects/{pid}/manuscript/sections/{K[kind]}/accept")

    m = api.post(f"/api/projects/{pid}/manuscript/revision", json={"letter": LETTER}).json()
    pts = m["revision"]["points"]
    assert [p["id"] for p in pts] == ["P1", "P2", "P3"]
    assert pts[0]["verbatim"] and not pts[2]["verbatim"]  # "overstates findings" is not word for word

    m = api.post(f"/api/projects/{pid}/manuscript/revision/points/P1/draft").json()
    p1 = m["revision"]["points"][0]
    assert p1["response"].startswith("We thank") and p1["change"]["section"] == K["methods"]
    call = claude.calls[-1]
    sys_text = call["system"] if isinstance(call["system"], str) else " ".join(b["text"] for b in call["system"])
    assert "Language of the response: English" in sys_text

    m = api.post(f"/api/projects/{pid}/manuscript/revision/points/P1/apply").json()
    p1 = m["revision"]["points"][0]
    assert p1["status"] == "done" and p1["applied"][0]["section"] == K["methods"]
    assert "clone YR145" in _sec(m, "methods")["source"]
    m = api.put(f"/api/projects/{pid}/manuscript/revision/points/P2", json={"response": "Added a comparison."}).json()
    assert m["revision"]["points"][1]["response"] == "Added a comparison."

    r = api.get(f"/api/projects/{pid}/manuscript/revision/export.zip")
    assert r.status_code == 200
    z = zipfile.ZipFile(io.BytesIO(r.content))
    import docx
    resp = "\n".join(p.text for p in docx.Document(io.BytesIO(z.read("response_to_reviewers.docx"))).paragraphs)
    assert "P1. Please provide the antibody clone" in resp and "Response: We thank" in resp and "Reviewer 2" in resp
    marked = docx.Document(io.BytesIO(z.read("manuscript_changes_marked.docx")))
    runs = [r for p in marked.paragraphs for r in p.runs]
    assert any(r.font.underline and "YR145" in r.text for r in runs)  # insertion marked
    assert any("Immunohistochemistry for CD117" in r.text and not r.font.underline for r in runs)


def test_letter_text_from_docx(api):
    import docx
    pid = api.post("/api/projects", json={"title": "x"}).json()["id"]
    d = docx.Document()
    d.add_paragraph("Reviewer 1: Please add the clone.")
    buf = io.BytesIO()
    d.save(buf)
    r = api.post(f"/api/projects/{pid}/manuscript/revision/letter-text", files={"file": ("decision.docx", buf.getvalue())})
    assert r.status_code == 200 and "Please add the clone" in r.json()["text"]
