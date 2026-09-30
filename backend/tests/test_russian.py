"""Russian-language journals: the manuscript is written in Russian, numbers in Russian notation, bilingual
title block, abstracts, captions and figures (e.g. «Архив патологии»)."""
import io
import zipfile

import pytest
from test_manuscript import Claude, _sec, project  # noqa: F401 — fixtures reused


@pytest.fixture()
def claude(monkeypatch):
    from app import llm
    c = Claude()
    monkeypatch.setattr(llm, "client", lambda: c)
    return c


@pytest.fixture()
def ru_project(api, project, claude):  # noqa: F811
    from app.journals import store as journals
    journals.seed()
    journals.set_fields("arkhiv-patologii", {
        "language": {"value": "ru"},
        "types.original_article.sections": {"value": ["Введение", "Материал и методы", "Результаты", "Обсуждение",
                                                      "Заключение"]},
        "types.original_article.abstract_structured": {"value": True},
        "types.original_article.abstract_headings": {"value": ["Цель исследования", "Материал и методы",
                                                               "Результаты", "Заключение"]},
        "types.original_article.abstract_words": {"value": 250},
        "statements.author_contributions": {"value": "required"},
        "statements.ai_disclosure": {"value": "required"},
    })
    pid, kit, expl = project
    r = api.put(f"/api/projects/{pid}/journal", json={"journal_id": "arkhiv-patologii"})
    assert r.status_code == 200, r.text
    return pid, kit, expl


def test_russian_notation_and_keys():
    from app.manuscript import lang as L
    from app.manuscript import store as ms
    assert L.fact("OR 0.0074, 95% CI 0.0003–0.21", "ru") == "ОШ 0,0074; 95% ДИ 0,0003–0,21"
    assert L.fact("median difference 2.00, 95% CI 1.00–2.00", "ru") == "разница медиан 2,00; 95% ДИ 1,00–2,00"
    assert L.fact("rho 0.81, 95% CI 0.57–0.93", "ru") == "ρ 0,81; 95% ДИ 0,57–0,93"
    assert L.fact("p < 0.001", "ru") == "p < 0,001" and L.fact("q = 0.020", "ru") == "q = 0,020"
    assert L.fact("OR 0.0074, 95% CI 0.0003–0.21", "en") == "OR 0.0074, 95% CI 0.0003–0.21"
    keys = [ms.slug(h) for h in ["Введение", "Материал и методы", "Результаты", "Обсуждение", "Заключение"]]
    assert keys == ["vvedenie", "material_i_metody", "rezultaty", "obsuzhdenie", "zaklyuchenie"]
    assert [ms.kind_of(h) for h in ["Введение", "Материал и методы", "Результаты", "Обсуждение", "Выводы",
                                    "Описание наблюдения"]] == [
        "introduction", "methods", "results", "discussion", "conclusion", "case_presentation"]


def test_russian_manuscript_flow(api, ru_project, claude):
    pid, kit, expl = ru_project
    api.post(f"/api/projects/{pid}/manuscript/terms/auto")
    m = api.get(f"/api/projects/{pid}/manuscript").json()
    assert m["lang"] == "ru" and m["figure_language"] == "ru"
    ids = {f["title"]: f["id"] for f in m["findings"]}
    a_kit, a_expl = ids[kit["title"]], ids[expl["title"]]

    api.post(f"/api/projects/{pid}/manuscript/plan")
    m = api.post(f"/api/projects/{pid}/manuscript/plan/approve").json()
    secs = m["sections"]
    assert [s["kind"] for s in secs] == ["abstract", "abstract_en", "introduction", "methods", "results",
                                         "discussion", "conclusion", "statements"]
    assert [s["heading"] for s in secs][:3] == ["Резюме", "Abstract", "Введение"]
    assert len({s["key"] for s in secs}) == len(secs)  # Russian headings get distinct keys
    K = {s["kind"]: s["key"] for s in secs}

    claude.sections = {
        "introduction": "ГИСО — наиболее частая мезенхимальная опухоль желудочно-кишечного тракта.",
        "methods": "Иммуногистохимическое исследование CD117 выполнено на срезах.",
        "results": (f"Серия включала {{{{n}}}} наблюдения ({{{{TAB:t1}}}}). Мутации KIT и PDGFRA "
                    f"взаимоисключались ({{{{{a_kit}.effect_expr}}}}; {{{{{a_kit}.p_expr}}}}). Это доказывает, что "
                    f"{{{{{a_expl}.p_expr}}}} верно."),
        "discussion": "Взаимоисключение мутаций наблюдалось в данной серии. Ограничения исследования: малая выборка.",
        "conclusion": "В данной серии мутации KIT и PDGFRA взаимоисключались.",
        "abstract": "### Цель исследования\n\nОписать {{n}} наблюдений ГИСО.",
        "abstract_en": "### Objective\n\nTo describe {{n}} GISTs; KIT and PDGFRA ({{" + a_kit + ".effect_expr}}).",
    }
    for kind in ("introduction", "methods", "results", "discussion", "conclusion"):
        r = api.post(f"/api/projects/{pid}/manuscript/sections/{K[kind]}/generate")
        assert r.status_code == 200, r.text
        m = r.json()
        sec = _sec(m, kind)
        codes = {i["code"] for i in sec["issues"]}
        assert "cyrillic" not in codes and "limitations" not in codes
        if kind == "results":
            rendered = "".join(s["v"] for p in sec["segments"] for s in p["segments"])
            assert "(табл. 1)" in rendered and "ОШ " in rendered and "95% ДИ" in rendered and "p " in rendered
            assert "95% CI" not in rendered and not any(ch.isdigit() and nxt == "." and after.isdigit()
                                                        for ch, nxt, after in zip(rendered, rendered[1:], rendered[2:]))
            assert "evidence_wording" in codes  # «доказывает» for an exploratory finding
        api.post(f"/api/projects/{pid}/manuscript/sections/{K[kind]}/accept")

    section_calls = [c for c in claude.calls if "(kind: " in c["messages"][0]["content"]]
    system = section_calls[-1]["system"]
    system = system if isinstance(system, str) else " ".join(b["text"] for b in system)
    assert "Статья пишется на русском языке" in system
    assert "Write the section in Russian" in section_calls[-1]["messages"][0]["content"]

    # the English abstract waits for the Russian one, keeps English notation and English rules
    assert api.post(f"/api/projects/{pid}/manuscript/sections/abstract_en/generate").status_code == 400
    api.post(f"/api/projects/{pid}/manuscript/sections/abstract/generate")
    api.post(f"/api/projects/{pid}/manuscript/sections/abstract/accept")
    m = api.post(f"/api/projects/{pid}/manuscript/sections/abstract_en/generate").json()
    en = _sec(m, "abstract_en")
    rendered = "".join(s["v"] for p in en["segments"] for s in p["segments"])
    assert "95% CI" in rendered and "ДИ" not in rendered
    last = claude.calls[-1]
    sys_en = last["system"] if isinstance(last["system"], str) else " ".join(b["text"] for b in last["system"])
    assert "Статья пишется на русском" not in sys_en and "Accepted Russian abstract" in last["messages"][0]["content"]
    api.post(f"/api/projects/{pid}/manuscript/sections/abstract_en/accept")

    inputs = {"ethics": {"committee": "Локальный этический комитет", "approval_number": "12/2025"},
              "consent": {"status": "waived"}, "coi": "Авторы заявляют об отсутствии конфликта интересов.",
              "funding": "Исследование не имело спонсорской поддержки.",
              "authors": [{"name": "Иванов А.С.", "roles": ["Conceptualization", "Writing – review & editing"]},
                          {"name": "Петров Н.Т.", "roles": ["Formal analysis", "Writing – original draft"]}]}
    api.put(f"/api/projects/{pid}/manuscript/inputs", json=inputs)
    m = api.post(f"/api/projects/{pid}/manuscript/sections/statements/generate").json()
    st = _sec(m, "statements")["source"]
    assert "### Участие авторов" in st and "Концепция и дизайн исследования — Иванов А.С." in st
    assert "Статистическая обработка данных — Петров Н.Т." in st
    assert "### Использование инструментов ИИ" in st and "При статистической обработке данных ИИ не использовался" in st
    api.post(f"/api/projects/{pid}/manuscript/sections/statements/accept")

    api.put(f"/api/projects/{pid}/manuscript/front", json={"title": "ГИСО: серия наблюдений",
                                                          "title_en": "GIST: a case series",
                                                          "keywords": ["ГИСО", "KIT"], "keywords_en": ["GIST", "KIT"]})
    fig = next(f for f in m["figures"] if f["kind"] == "finding")
    api.put(f"/api/projects/{pid}/manuscript/assets",
            json={"figures": [{"id": fig["id"], "include": True, "caption": "", "caption_en": "My English caption"}]})
    r = api.post(f"/api/projects/{pid}/manuscript/export", params={"mode": "draft"})
    assert r.status_code == 200, r.text
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = set(z.namelist())
    assert {"figures/Рис-1.png", "figures/Рис-1_en.png"} <= names
    import docx
    doc = docx.Document(io.BytesIO(z.read("manuscript.docx")))
    text = "\n".join(p.text for p in doc.paragraphs)
    for s in ("ГИСО: серия наблюдений", "GIST: a case series", "Ключевые слова: ГИСО; KIT", "Keywords: GIST; KIT",
              "Резюме", "Abstract", "Литература/References", "Таблица 1.", "Table 1.", "Подписи к рисункам",
              "Рис. 1.", "Fig. 1. My English caption", "Дополнительная информация"):
        assert s in text, s
    assert doc.tables[0].rows[0].cells[0].text == "Признак"
    assert doc.tables[0].rows[0].cells[1].text == "Все случаи (n = 24)"
    status = api.get(f"/api/projects/{pid}/manuscript/export/status").json()
    assert not any(b["code"] in ("cyrillic_table", "cyrillic_figure", "cyrillic") for b in status["blocking"])


def test_russian_cover_letter_frame():
    from app.manuscript import cover
    blocks = cover.assemble({"settings": {"corresponding_name": "Иванов А.С."}}, "Направляем рукопись.",
                            "Архив патологии", "ru")
    text = cover.as_text(blocks)
    assert "Главному редактору журнала «Архив патологии»" in text and "Уважаемый главный редактор!" in text
    assert "С уважением," in text and "Dear" not in text
    assert "Write the letter in Russian" in cover.SYSTEM_RU
