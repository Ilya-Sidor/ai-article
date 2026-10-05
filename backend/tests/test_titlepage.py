"""Title page: authors with affiliation superscripts, ORCID, bilingual block for Russian journals."""
import io


def test_translit_and_orcid():
    from app.manuscript import titlepage as T
    assert T.translit("Белая Жанна Евгеньевна") == "Belaia Zhanna Evgen'evna"
    assert T.translit("Мельниченко") == "Mel'nichenko" and T.translit("Щукин Юрий") == "Shchukin Iurii"
    assert T.orcid_ok("0000-0002-1825-0097") and not T.orcid_ok("0000-0002-1825-0098")
    assert not T.orcid_ok("1234")


def test_issues_and_blocks():
    from app.manuscript import titlepage as T
    inputs = {"affiliations": [{"name": "ФГБУ «НМИЦ онкологии»", "address": "Москва, Россия"}],
              "authors": [{"name": "Иванов Иван Иванович", "affiliations": [1], "orcid": "0000-0002-1825-0098"},
                          {"name": "Петров Пётр", "affiliations": [2], "corresponding": True}]}
    codes = {i["code"]: i for i in T.issues(inputs, "ru")}
    assert codes["titlepage_corresponding"]["severity"] == "blocking"  # no e-mail
    assert "titlepage_orcid" in codes and "titlepage_aff_en" in codes and "titlepage_author_aff" in codes
    inputs["authors"][1]["email"] = "p@example.org"
    blocks = T.blocks(inputs, {"title": "ГИСО", "title_en": "GIST"}, "ru")
    kinds = [k for k, _ in blocks]
    assert kinds.count("title") == 2 and kinds.count("authors") == 2 and "break" in kinds
    en_authors = [p for k, p in blocks if k == "authors"][1]
    assert en_authors[0] == ("Ivanov Ivan Ivanovich", "1")


def test_export_title_page(api):
    import docx

    from app.manuscript import export
    inputs = {"affiliations": [{"name": "Hospital A", "name_en": "Hospital A", "address": "City"}],
              "authors": [{"name": "Anna Smith", "affiliations": [1], "email": "a@x.org", "corresponding": True}]}

    class R:
        bibliography = []

        def segments(self, source, kind=None):
            return []
    ctx = {"renderer": R(), "front": {"title": "T"}, "inputs": inputs, "sections": [], "counts": {}, "tables": [],
           "figures": [], "lang": "en"}
    doc = docx.Document(io.BytesIO(export.build_docx(ctx, False, [])))
    p = doc.paragraphs[1]
    assert p.runs[0].text == "Anna Smith" and p.runs[1].text == "1" and p.runs[1].font.superscript
    text = "\n".join(x.text for x in doc.paragraphs)
    assert "Hospital A, City" in text and "Corresponding author: Anna Smith, a@x.org" in text
